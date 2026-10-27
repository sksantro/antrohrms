"""Excel import — synthetic workbooks only (BRD §20.1: no real employee data in tests)."""
from datetime import date, datetime
from io import BytesIO
from unittest import mock

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl.styles import PatternFill

from accounts.models import User
from common.models import AuditEvent
from common.testing import Fixture
from employees import importer
from employees.models import Department, Employee

HEAD = ["Emp-Id", "First Name", "Last Name", "Date of Birth", "Aadhaar number", "Pan Number", "Personal Email Address",
        "Company Email Address", "Contact Number", "Address", "Date of Hire", "Designation", "Department",
        "Reporting Manager", "Employment Status", "Bank Account Details", "Emergency Contact Name", "Emergency Contact Number"]
TODAY = date(2026, 10, 27)
GREY = PatternFill("solid", fgColor="FFCCCCCC")


def book(rows, grey_rows=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEAD)
    for i, r in enumerate(rows, start=2):
        ws.append([r.get(h) for h in HEAD])
        if i in grey_rows:
            for c in ws[i]:
                c.fill = GREY
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def person(**kw):
    base = {"Emp-Id": "APT0101", "First Name": "test", "Last Name": "person", "Date of Birth": datetime(1999, 1, 1),
            "Aadhaar number": 123412341234.0, "Pan Number": "ABCDE1234F", "Personal Email Address": "tp@example.com",
            "Company Email Address": None, "Contact Number": 9000000001.0, "Address": "Street 1", "Date of Hire": datetime(2025, 1, 6),
            "Designation": "software engineer", "Department": "IT", "Reporting Manager": "", "Employment Status": "Full-Time",
            "Bank Account Details": "000111222", "Emergency Contact Name": "kin one", "Emergency Contact Number": 9000000002.0}
    base.update(kw)
    return base


class ParseTests(Fixture):
    def parse(self, rows, grey=()):
        return importer.parse(book(rows, grey), self.org, today=TODAY)

    def test_maps_columns_and_ignores_sensitive_ones(self):
        res = self.parse([person()])
        r = res["rows"][0]
        self.assertEqual((r.emp_id, r.name, r.login_email, r.email_source), ("APT0101", "Test Person", "tp@example.com", "PERSONAL"))
        self.assertEqual((r.status, r.employment_type, r.department, r.phone), ("ACTIVE_CONFIRMED", "FULL_TIME", "IT", "9000000001"))
        self.assertEqual(r.emergency, {"name": "Kin One", "phone": "9000000002"})
        for col in ("Aadhaar number", "Pan Number", "Bank Account Details", "Date of Birth"):
            self.assertIn(col, res["ignored_columns"])
        self.assertNotIn("Company Email Address", res["ignored_columns"])
        self.assertEqual(res["new_departments"], ["IT"])

    def test_company_email_preferred_and_typos_flagged(self):
        r = self.parse([person(**{"Company Email Address": "Tp@Antro.example"})])["rows"][0]
        self.assertEqual((r.login_email, r.email_source), ("tp@antro.example", "COMPANY"))
        typo = self.parse([person(**{"Personal Email Address": "a@gmial.com"})])["rows"][0]
        self.assertTrue(any("typo" in w for w in typo.warnings))

    def test_status_from_joining_date_and_text_dates(self):
        rows = self.parse([
            person(**{"Emp-Id": "APT0102", "Personal Email Address": "a@example.com", "Date of Hire": "14-08-2026"}),
            person(**{"Emp-Id": "APT0103", "Personal Email Address": "b@example.com", "Date of Hire": datetime(2026, 12, 10)}),
            person(**{"Emp-Id": "APT0104", "Personal Email Address": "c@example.com", "Date of Hire": "31/02/2026"}),
        ])["rows"]
        self.assertEqual(rows[0].status, "PROBATION")  # ends 14 Nov 2026
        self.assertEqual(rows[1].status, "PREBOARDING")
        self.assertTrue(any("swapped" in w for w in rows[1].warnings))
        self.assertEqual(rows[2].action, "SKIP")

    def test_grey_rows_are_former_employees_even_without_email(self):
        rows = self.parse([person(), person(**{"Emp-Id": "APT0105", "Personal Email Address": None})], grey=(3,))["rows"]
        self.assertEqual((rows[1].action, rows[1].status), ("CREATE_FORMER", "EXITED"))
        self.assertTrue(rows[1].login_email.endswith("@no-login.invalid"))

    def test_rows_that_cannot_be_imported(self):
        rows = self.parse([
            person(**{"Personal Email Address": None}),                                   # current, no email
            person(**{"Emp-Id": "APT0106", "Personal Email Address": "emp@test.example"}),  # already a user
            person(**{"Emp-Id": "T-emp", "Personal Email Address": "z@example.com"}),         # code already used
            person(**{"Emp-Id": "APT0107", "Personal Email Address": "dup@example.com"}),
            person(**{"Emp-Id": "APT0108", "Personal Email Address": "dup@example.com"}),
        ])["rows"]
        self.assertEqual([r.action for r in rows], ["SKIP", "SKIP", "SKIP", "CREATE", "SKIP"])

    def test_manager_matched_by_first_and_last_name(self):
        rows = self.parse([person(**{"Reporting Manager": "mgr (team lead)"}), person(**{"Emp-Id": "APT0109", "Personal Email Address": "x@example.com", "Reporting Manager": "Nobody Here"})])["rows"]
        self.assertEqual(rows[0].manager_name, "Mgr")
        self.assertTrue(any("isn't in the system" in w for w in rows[1].warnings))

    def test_rejects_non_workbook_and_missing_columns(self):
        with self.assertRaises(ValueError):
            importer.parse(b"not a workbook", self.org)
        wb = openpyxl.Workbook(); wb.active.append(["Name", "Phone"]); buf = BytesIO(); wb.save(buf)
        with self.assertRaisesMessage(ValueError, "Missing required column"):
            importer.parse(buf.getvalue(), self.org)


class ImportApiTests(Fixture):
    def upload(self, user, data, commit=False, name="team.xlsx"):
        f = SimpleUploadedFile(name, data, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        return self.as_user(user).post("/api/v1/admin/employees/import", {"file": f, "commit": "true" if commit else "false"}, format="multipart")

    def test_only_super_admin(self):
        data = book([person()])
        for u in (self.hr, self.emp):
            self.assertEqual(self.upload(u, data).status_code, 403)

    def test_preview_writes_nothing(self):
        r = self.upload(self.admin, book([person()]))
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["summary"]["create_current"], 1)
        self.assertFalse(Employee.objects.filter(employee_code="APT0101").exists())
        self.assertNotIn("123412341234", r.content.decode())  # Aadhaar never echoed back

    @mock.patch("employees.importer.timezone.localdate", return_value=TODAY)
    def test_commit_creates_current_and_former_with_invites(self, _):
        data = book([
            person(**{"Reporting Manager": "Mgr"}),
            person(**{"Emp-Id": "APT0120", "First Name": "Old", "Last Name": "Timer", "Personal Email Address": "old@example.com", "Department": None}),
        ], grey_rows=(3,))
        r = self.upload(self.admin, data, commit=True)
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual([i["employee_code"] for i in r.data["invites"]], ["APT0101"])
        self.assertIn("/activate?token=", r.data["invites"][0]["invite_url"])
        cur = Employee.objects.get(employee_code="APT0101")
        self.assertEqual((cur.status, cur.reporting_manager, cur.department.name, cur.user.is_active), ("ACTIVE_CONFIRMED", self.mgr.employee, "IT", True))
        old = Employee.objects.get(employee_code="APT0120")
        self.assertEqual((old.status, old.user.is_active, old.department.name), ("EXITED", False, "Unassigned"))
        self.assertFalse(old.user.has_usable_password())
        self.assertTrue(AuditEvent.objects.filter(action="EMPLOYEES_BULK_IMPORTED").exists())
        # next generated code continues the sheet's numbering
        from employees.services import issue_employee_code
        self.assertEqual(issue_employee_code(self.org), "APT0121")

    def test_former_employee_cannot_log_in(self):
        self.upload(self.admin, book([person(**{"Emp-Id": "APT0130"})], grey_rows=(2,)), commit=True)
        u = User.objects.get(email="tp@example.com")
        u.set_password("Whatever-123456")
        u.save()
        from rest_framework.test import APIClient
        c = APIClient()
        self.assertEqual(c.post("/api/v1/auth/login", {"email": "tp@example.com", "password": "Whatever-123456"}, format="json").status_code, 422)

    def test_rejects_wrong_file_type(self):
        r = self.upload(self.admin, b"a,b", name="team.csv")
        self.assertEqual(r.data["code"], "INVALID_FILE")

    def test_reimport_skips_existing_rows(self):
        data = book([person()])
        self.upload(self.admin, data, commit=True)
        r = self.upload(self.admin, data)
        self.assertEqual(r.data["summary"]["skipped"], 1)
        self.assertEqual(Department.objects.filter(name="IT").count(), 1)
