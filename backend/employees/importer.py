"""Bulk employee import from an Excel sheet (BRD §27.1: staged mapping, validation preview, reject bad rows).

Two passes over the same upload: `preview` (no writes) and `commit` (one transaction).
Highly sensitive columns (Aadhaar, PAN, bank, tax ID, date of birth, gender) are never read into the
system; they are listed back as ignored so the person importing knows they weren't stored.
Grey-highlighted rows are treated as former employees: kept as EXITED records, inactive, no login.
"""
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO

import openpyxl
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.utils import timezone

from accounts.models import Role, RoleAssignment, User
from common.services import audit

from .models import Department, Employee, EmployeeCodeSequence, add_months
from .services import send_invite

MAX_ROWS = 1000
UNASSIGNED = "Unassigned"

FIELDS = {
    "emp_id": ["emp-id", "emp id", "employee id", "employee code", "emp code", "code"],
    "first_name": ["first name"],
    "last_name": ["last name", "surname"],
    "full_name": ["name", "full name", "employee name"],
    "phone": ["contact number", "phone", "phone number", "mobile", "mobile number"],
    "personal_email": ["personal email address", "personal email"],
    "company_email": ["company email address", "company email", "work email", "official email"],
    "address": ["address"],
    "start_date": ["date of hire", "date of joining", "joining date", "doj", "start date"],
    "designation": ["designation"],
    "job_title": ["job title"],
    "department": ["department"],
    "manager": ["reporting manager", "manager"],
    "employment_type": ["employment status", "employment type"],
    "ec_name": ["emergency contact name"],
    "ec_relation": ["emergency contact relationship"],
    "ec_phone": ["emergency contact number"],
}
SENSITIVE = re.compile(r"\b(aadhaar|aadhar|pan|bank|tax|date of birth|dob|gender|salary|ctc)\b")
TYPO_DOMAINS = {"gmial.com", "gmal.com", "gmai.com", "gamil.com", "gmail.co", "gnail.com", "yaho.com", "hotmial.com"}


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def _key(s) -> str:
    return _norm(s).lower()


@dataclass
class Row:
    sheet_row: int
    emp_id: str = ""
    name: str = ""
    login_email: str = ""
    email_source: str = ""  # COMPANY | PERSONAL | NONE
    former: bool = False
    status: str = ""
    start_date: date | None = None
    department: str = ""
    designation: str = ""
    employment_type: str = "FULL_TIME"
    manager_text: str = ""
    manager_id: str | None = None
    manager_name: str = ""
    phone: str = ""
    address: str = ""
    emergency: dict = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def action(self) -> str:
        return "SKIP" if self.errors else ("CREATE_FORMER" if self.former else "CREATE")

    def as_dict(self) -> dict:
        return {
            "sheet_row": self.sheet_row, "emp_id": self.emp_id, "name": self.name, "login_email": self.login_email,
            "email_source": self.email_source, "former": self.former, "status": self.status,
            "start_date": self.start_date, "department": self.department, "designation": self.designation,
            "employment_type": self.employment_type, "manager_text": self.manager_text, "manager_name": self.manager_name,
            "action": self.action, "errors": self.errors, "warnings": self.warnings,
        }


def _is_grey(cells) -> bool:
    """Most filled cells share a neutral grey fill (R≈G≈B, not white)."""
    greys = 0
    filled = [c for c in cells if c.value not in (None, "")]
    for c in filled:
        f = c.fill
        rgb = getattr(f.fgColor, "rgb", None) if f is not None and f.fill_type == "solid" else None
        if isinstance(rgb, str) and len(rgb) == 8:
            r, g, b = (int(rgb[i:i + 2], 16) for i in (2, 4, 6))
            if max(r, g, b) - min(r, g, b) <= 12 and max(r, g, b) < 0xF0:
                greys += 1
    return bool(filled) and greys >= len(filled) / 2


def _phone(v) -> str:
    if v in (None, ""):
        return ""
    if isinstance(v, float):
        v = int(v)
    digits = re.sub(r"\D", "", str(v))
    return digits[-12:]


def _date(v, row: Row) -> date | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _norm(v)
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d", "%d %b %Y", "%d %B %Y", "%d-%b-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    row.errors.append(f"Couldn't read the joining date “{s}” (use DD-MM-YYYY).")
    return None


def _title(s: str) -> str:
    s = _norm(s)
    return s.title() if s and (s.islower() or s.isupper()) else s


def _dept_name(s: str) -> str:
    s = _norm(s)
    if not s:
        return ""
    return s.upper() if s.replace("/", "").isalpha() and len(s.replace("/", "")) <= 4 else (s[0].upper() + s[1:])


def _email(v) -> str:
    s = _norm(v).lower()
    if not s or s in {"-", "na", "n/a", "nil", "none"}:
        return ""
    return s


def _valid_email(s: str) -> bool:
    try:
        validate_email(s)
        return True
    except ValidationError:
        return False


def _name_tokens(s: str) -> list[str]:
    s = re.sub(r"\([^)]*\)", " ", s)  # drop notes like "(team lead)"
    return [t for t in re.split(r"[^a-z]+", s.lower()) if t]


def parse(data: bytes, org, today: date | None = None) -> dict:
    today = today or timezone.localdate()
    try:
        wb = openpyxl.load_workbook(BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001 — any parse failure means "not a readable .xlsx"
        raise ValueError("That file couldn't be read as an Excel (.xlsx) workbook.") from exc
    ws = wb.worksheets[0]
    if ws.max_row > 20 * MAX_ROWS:  # formatted-but-empty rows are common; only a huge sheet is refused up front
        raise ValueError(f"The sheet is too large. Keep it under {MAX_ROWS} employees.")

    header = [_key(c.value) for c in ws[1]]
    cols: dict[str, int] = {}
    for name, aliases in FIELDS.items():
        for i, h in enumerate(header):
            if h in aliases and name not in cols:
                cols[name] = i
    ignored = [_norm(ws.cell(1, i + 1).value) for i, h in enumerate(header) if h and i not in cols.values() and SENSITIVE.search(h)]
    unmapped = [_norm(ws.cell(1, i + 1).value) for i, h in enumerate(header)
                if h and i not in cols.values() and _norm(ws.cell(1, i + 1).value) not in ignored]
    missing = [n for n in ("emp_id", "start_date") if n not in cols]
    if "first_name" not in cols and "full_name" not in cols:
        missing.append("name")
    if "company_email" not in cols and "personal_email" not in cols:
        missing.append("email")
    if missing:
        raise ValueError(f"Missing required column(s): {', '.join(missing)}. Expected headers like “Emp-Id”, “First Name”, “Date of Hire”, “Company Email Address”.")

    existing_emails = set(User.objects.values_list("email", flat=True))
    existing_codes = {c.upper() for c in Employee.objects.filter(organization=org).exclude(employee_code=None).values_list("employee_code", flat=True)}
    people = list(Employee.objects.filter(organization=org).values("id", "display_name"))
    dept_lookup = {d.name.lower(): d.name for d in Department.objects.filter(organization=org)}

    rows: list[Row] = []
    seen_emails: dict[str, int] = {}
    seen_codes: dict[str, int] = {}
    for cells in ws.iter_rows(min_row=2):
        cells = list(cells)
        if not any(c.value not in (None, "") for c in cells):
            continue
        if len(rows) >= MAX_ROWS:
            raise ValueError(f"The sheet has more than {MAX_ROWS} employees. Split it into smaller files.")
        get = lambda n: cells[cols[n]].value if n in cols and cols[n] < len(cells) else None  # noqa: E731
        r = Row(sheet_row=cells[0].row, former=_is_grey(cells))
        r.emp_id = _norm(get("emp_id")).upper()
        full = _norm(get("full_name")) or f"{_norm(get('first_name'))} {_norm(get('last_name'))}"
        r.name = _title(full)
        company, personal = _email(get("company_email")), _email(get("personal_email"))
        if company and _valid_email(company):
            r.login_email, r.email_source = company, "COMPANY"
        elif personal and _valid_email(personal):
            r.login_email, r.email_source = personal, "PERSONAL"
            if not r.former:
                r.warnings.append("No company email, so their personal email will be their login.")
        if company and not _valid_email(company):
            r.warnings.append(f"Company email “{company}” isn't valid and was ignored.")
        domain = r.login_email.split("@")[-1] if r.login_email else ""
        if domain in TYPO_DOMAINS or (domain.endswith("gmail.com") and domain != "gmail.com"):
            r.warnings.append(f"“{domain}” looks like a typo. Fix it in the sheet before importing.")
        r.start_date = _date(get("start_date"), r)
        r.designation = _title(get("designation")) or _title(get("job_title"))
        dept = _dept_name(get("department"))
        r.department = dept_lookup.get(dept.lower(), dept) if dept else UNASSIGNED
        if not dept and not r.former:
            r.warnings.append(f"No department; will go to “{UNASSIGNED}”.")
        et = _key(get("employment_type"))
        r.employment_type = "INTERN" if "intern" in et else "CONTRACT" if "contract" in et else "FULL_TIME"
        r.phone = _phone(get("phone"))
        r.address = _norm(get("address"))[:1000]
        ec = {"name": _title(get("ec_name")), "relationship": _title(get("ec_relation")), "phone": _phone(get("ec_phone"))}
        r.emergency = {k: v for k, v in ec.items() if v}
        r.manager_text = _norm(get("manager"))

        # --- validation ---
        if not r.emp_id:
            r.errors.append("Emp-Id is empty.")
        if not r.name.strip():
            r.errors.append("Name is empty.")
        if not r.start_date and not r.errors:
            r.errors.append("Joining date is empty.")
        if not r.login_email:
            if r.former and r.emp_id:
                r.login_email, r.email_source = f"former.{r.emp_id.lower()}@no-login.invalid", "NONE"
            else:
                r.errors.append("No valid email, so an account can't be created.")
        if r.emp_id in existing_codes:
            r.errors.append(f"{r.emp_id} is already in the system.")
        if r.login_email in existing_emails:
            r.errors.append(f"{r.login_email} is already in the system.")
        if r.emp_id and r.emp_id in seen_codes:
            r.errors.append(f"Emp-Id repeats sheet row {seen_codes[r.emp_id]}.")
        if r.login_email and r.login_email in seen_emails:
            r.errors.append(f"Email repeats sheet row {seen_emails[r.login_email]}.")
        seen_codes.setdefault(r.emp_id, r.sheet_row)
        seen_emails.setdefault(r.login_email, r.sheet_row)
        if not r.designation:
            r.designation = "Not specified"
            if not r.former:
                r.warnings.append("No designation.")

        # --- status ---
        if r.former:
            r.status = Employee.Status.EXITED
        elif r.start_date and r.start_date > today:
            r.status = Employee.Status.PREBOARDING
            if r.start_date.day <= 12:
                r.warnings.append(f"Joins in the future ({r.start_date:%d %b %Y}). If that's wrong, day and month may be swapped in Excel.")
        elif r.start_date and add_months(r.start_date, 3) <= today:
            r.status = Employee.Status.ACTIVE_CONFIRMED
        else:
            r.status = Employee.Status.PROBATION
        rows.append(r)

    # --- managers: match existing people or current people in this sheet by first + last name ---
    batch = [(f"row:{r.sheet_row}", r.name) for r in rows if not r.errors and not r.former]
    candidates = [(str(p["id"]), p["display_name"]) for p in people] + batch
    for r in rows:
        if not r.manager_text:
            continue
        first_choice = r.manager_text.split(",")[0]
        toks = _name_tokens(first_choice)
        hits = [(cid, n) for cid, n in candidates if toks and _name_tokens(n)[:1] == toks[:1] and _name_tokens(n)[-1:] == toks[-1:]]
        if len(hits) == 1:
            r.manager_id, r.manager_name = hits[0]
        elif not r.former:
            r.warnings.append(f"Manager “{first_choice}” isn't in the system yet. Set it later in People.")

    creates = [r for r in rows if r.action != "SKIP"]
    new_depts = sorted({r.department for r in creates if r.department.lower() not in dept_lookup})
    return {
        "rows": rows,
        "ignored_columns": ignored,
        "unmapped_columns": unmapped,
        "new_departments": new_depts,
        "summary": {
            "total": len(rows),
            "create_current": sum(r.action == "CREATE" for r in rows),
            "create_former": sum(r.action == "CREATE_FORMER" for r in rows),
            "skipped": sum(r.action == "SKIP" for r in rows),
            "with_warnings": sum(bool(r.warnings) for r in creates),
        },
    }


def _bump_sequence(org) -> None:
    """Continue the sheet's numbering for new hires: next code = highest APT number + 1."""
    seq, _ = EmployeeCodeSequence.objects.select_for_update().get_or_create(organization=org)
    nums = [int(m.group(1)) for c in Employee.objects.filter(organization=org).exclude(employee_code=None).values_list("employee_code", flat=True)
            if (m := re.fullmatch(rf"{re.escape(seq.prefix)}-?(\d+)", c))]
    if nums and max(nums) + 1 > seq.next_value:
        seq.next_value = max(nums) + 1
        seq.save(update_fields=["next_value"])


@transaction.atomic
def commit(data: bytes, *, actor, request=None) -> dict:
    org = actor.organization
    result = parse(data, org)
    rows = [r for r in result["rows"] if r.action != "SKIP"]
    depts = {d.name.lower(): d for d in Department.objects.filter(organization=org)}
    for name in result["new_departments"]:
        depts[name.lower()] = Department.objects.create(organization=org, name=name)
    created: dict[str, Employee] = {}
    invites = []
    for r in rows:
        user = User.objects.create_user(email=r.login_email, display_name=r.name, organization=org, is_active=not r.former)
        RoleAssignment.objects.create(user=user, role=Role.EMPLOYEE, granted_by=actor)
        emp = Employee.objects.create(
            organization=org, user=user, employee_code=r.emp_id, display_name=r.name, official_email=r.login_email,
            phone=r.phone, department=depts[r.department.lower()], designation=r.designation[:120],
            employment_type=r.employment_type, start_date=r.start_date, status=r.status,
            address=r.address, emergency_contact=r.emergency,
        )
        created[f"row:{r.sheet_row}"] = emp
        audit(actor=actor, action="EMPLOYEE_IMPORTED", entity=emp, after={"code": emp.employee_code, "status": emp.status, "sheet_row": r.sheet_row}, request=request)
        if not r.former:
            _, link = send_invite(user=user, actor=actor, request=request)
            invites.append({"name": r.name, "employee_code": r.emp_id, "login_email": r.login_email, "invite_url": link})
    for r in rows:
        if r.manager_id:
            mgr_id = created[r.manager_id].pk if r.manager_id.startswith("row:") else r.manager_id
            Employee.objects.filter(pk=created[f"row:{r.sheet_row}"].pk).update(reporting_manager_id=mgr_id)
    _bump_sequence(org)
    audit(actor=actor, action="EMPLOYEES_BULK_IMPORTED", entity=actor, after=result["summary"], request=request, organization=org)
    return {"summary": result["summary"], "invites": invites, "new_departments": result["new_departments"]}
