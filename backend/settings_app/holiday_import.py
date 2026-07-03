import csv
import io
import re
from datetime import date, datetime
from typing import BinaryIO

from django.core.files.uploadedfile import UploadedFile

from settings_app.models import CompanyHoliday

DATE_PATTERNS = [
    ('%Y-%m-%d', re.compile(r'\b(\d{4}-\d{2}-\d{2})\b')),
    ('%d/%m/%Y', re.compile(r'\b(\d{1,2}/\d{1,2}/\d{4})\b')),
    ('%d-%m-%Y', re.compile(r'\b(\d{1,2}-\d{1,2}-\d{4})\b')),
    ('%d.%m.%Y', re.compile(r'\b(\d{1,2}\.\d{1,2}\.\d{4})\b')),
]

HEADER_ALIASES = {
    'name': {'name', 'holiday', 'holiday name', 'holiday_name', 'title'},
    'date': {'date', 'holiday date', 'holiday_date', 'day'},
    'holiday_type': {'type', 'holiday type', 'holiday_type', 'category'},
    'description': {'description', 'desc', 'details', 'remarks', 'note'},
}

HOLIDAY_TYPE_ALIASES = {
    'company_holiday': CompanyHoliday.HolidayType.COMPANY_HOLIDAY,
    'company holiday': CompanyHoliday.HolidayType.COMPANY_HOLIDAY,
    'company': CompanyHoliday.HolidayType.COMPANY_HOLIDAY,
    'optional_holiday': CompanyHoliday.HolidayType.OPTIONAL_HOLIDAY,
    'optional holiday': CompanyHoliday.HolidayType.OPTIONAL_HOLIDAY,
    'optional': CompanyHoliday.HolidayType.OPTIONAL_HOLIDAY,
    'festival': CompanyHoliday.HolidayType.FESTIVAL,
    'national_holiday': CompanyHoliday.HolidayType.NATIONAL_HOLIDAY,
    'national holiday': CompanyHoliday.HolidayType.NATIONAL_HOLIDAY,
    'national': CompanyHoliday.HolidayType.NATIONAL_HOLIDAY,
}


def _normalize_header(value: str) -> str:
    return re.sub(r'\s+', ' ', (value or '').strip().lower())


def _map_headers(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, header in enumerate(headers):
        normalized = _normalize_header(header)
        for field, aliases in HEADER_ALIASES.items():
            if normalized in aliases and field not in mapping:
                mapping[field] = index
    return mapping


def parse_date_value(raw_value) -> date | None:
    if raw_value is None:
        return None
    if isinstance(raw_value, datetime):
        return raw_value.date()
    if isinstance(raw_value, date):
        return raw_value

    text = str(raw_value).strip()
    if not text:
        return None

    for fmt, pattern in DATE_PATTERNS:
        match = pattern.search(text)
        if not match:
            continue
        try:
            return datetime.strptime(match.group(1), fmt).date()
        except ValueError:
            continue
    return None


def normalize_holiday_type(raw_value) -> str:
    if not raw_value:
        return CompanyHoliday.HolidayType.COMPANY_HOLIDAY
    text = str(raw_value).strip()
    if text in CompanyHoliday.HolidayType.values:
        return text
    alias = HOLIDAY_TYPE_ALIASES.get(_normalize_header(text))
    return alias or CompanyHoliday.HolidayType.COMPANY_HOLIDAY


def _build_preview_row(
    row_number: int,
    name: str,
    parsed_date: date | None,
    holiday_type: str,
    description: str,
) -> dict:
    errors: list[str] = []
    clean_name = (name or '').strip()
    if not clean_name:
        errors.append('Holiday name is required.')
    if not parsed_date:
        errors.append('Valid date is required.')

    return {
        'row_number': row_number,
        'name': clean_name,
        'date': parsed_date.isoformat() if parsed_date else '',
        'holiday_type': holiday_type,
        'description': (description or '').strip(),
        'is_valid': not errors,
        'errors': errors,
    }


def _rows_from_mapping(rows: list[list], header_map: dict[str, int], start_row: int = 2) -> list[dict]:
    preview_rows: list[dict] = []
    for offset, row in enumerate(rows, start=start_row):
        if not row or not any(str(cell).strip() for cell in row):
            continue
        name = row[header_map['name']] if 'name' in header_map and header_map['name'] < len(row) else ''
        raw_date = row[header_map['date']] if 'date' in header_map and header_map['date'] < len(row) else ''
        raw_type = (
            row[header_map['holiday_type']]
            if 'holiday_type' in header_map and header_map['holiday_type'] < len(row)
            else ''
        )
        description = (
            row[header_map['description']]
            if 'description' in header_map and header_map['description'] < len(row)
            else ''
        )
        preview_rows.append(
            _build_preview_row(
                offset,
                str(name),
                parse_date_value(raw_date),
                normalize_holiday_type(raw_type),
                str(description),
            ),
        )
    return preview_rows


def parse_csv_file(file_obj: BinaryIO) -> tuple[list[dict], list[str]]:
    warnings: list[str] = []
    content = file_obj.read()
    if isinstance(content, bytes):
        text = content.decode('utf-8-sig')
    else:
        text = content

    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return [], ['The uploaded CSV file is empty.']

    header_map = _map_headers(rows[0])
    if 'name' not in header_map or 'date' not in header_map:
        return [], [
            'CSV must include headers for holiday name and date '
            '(e.g. name, date, holiday_type, description).',
        ]

    return _rows_from_mapping(rows[1:], header_map), warnings


def parse_excel_file(file_obj: BinaryIO) -> tuple[list[dict], list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError('Excel support is not available on the server.') from exc

    warnings: list[str] = []
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    sheet = workbook.active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()

    if not rows:
        return [], ['The uploaded Excel file is empty.']

    headers = [str(cell or '') for cell in rows[0]]
    header_map = _map_headers(headers)
    if 'name' not in header_map or 'date' not in header_map:
        return [], [
            'Excel must include headers for holiday name and date '
            '(e.g. name, date, holiday_type, description).',
        ]

    return _rows_from_mapping(rows[1:], header_map), warnings


def _parse_text_lines(text: str) -> list[dict]:
    preview_rows: list[dict] = []
    row_number = 0
    for line in text.splitlines():
        clean_line = line.strip()
        if not clean_line or len(clean_line) < 4:
            continue
        parsed_date = None
        matched_text = ''
        for fmt, pattern in DATE_PATTERNS:
            match = pattern.search(clean_line)
            if not match:
                continue
            try:
                parsed_date = datetime.strptime(match.group(1), fmt).date()
                matched_text = match.group(1)
                break
            except ValueError:
                continue
        if not parsed_date:
            continue

        row_number += 1
        name = clean_line.replace(matched_text, '').strip(' -–:|,')
        if not name:
            name = f'Holiday {row_number}'

        preview_rows.append(
            _build_preview_row(
                row_number,
                name,
                parsed_date,
                CompanyHoliday.HolidayType.COMPANY_HOLIDAY,
                'Imported from image',
            ),
        )
    return preview_rows


def parse_image_file(file_obj: BinaryIO) -> tuple[list[dict], list[str]]:
    warnings: list[str] = []
    try:
        from PIL import Image
    except ImportError as exc:
        raise ValueError('Image processing is not available on the server.') from exc

    try:
        import pytesseract
    except ImportError as exc:
        raise ValueError(
            'Image OCR is not available. Please upload CSV or Excel, or install OCR support on the server.',
        ) from exc

    image = Image.open(file_obj)
    try:
        text = pytesseract.image_to_string(image)
    except Exception as exc:
        raise ValueError(
            'Unable to read text from the image. Try a clearer image or upload CSV/Excel instead.',
        ) from exc

    rows = _parse_text_lines(text)
    if not rows:
        warnings.append(
            'No holiday rows could be detected in the image. Try CSV/Excel for best results.',
        )
    else:
        warnings.append(
            'Image rows were auto-detected. Please review names and dates before importing.',
        )
    return rows, warnings


def parse_holiday_upload(uploaded_file: UploadedFile) -> tuple[list[dict], list[str], str]:
    filename = (uploaded_file.name or '').lower()
    file_obj = uploaded_file.file
    file_obj.seek(0)

    if filename.endswith('.csv'):
        rows, warnings = parse_csv_file(file_obj)
        return rows, warnings, 'csv'
    if filename.endswith(('.xlsx', '.xlsm', '.xltx', '.xltm')):
        rows, warnings = parse_excel_file(file_obj)
        return rows, warnings, 'excel'
    if filename.endswith(('.png', '.jpg', '.jpeg', '.webp', '.bmp', '.gif', '.tif', '.tiff')):
        rows, warnings = parse_image_file(file_obj)
        return rows, warnings, 'image'

    raise ValueError('Unsupported file type. Upload CSV, Excel (.xlsx), or an image (PNG/JPG).')


def import_holiday_rows(items: list[dict]) -> dict:
    created = 0
    updated = 0
    skipped = 0
    errors: list[str] = []

    for item in items:
        if not item.get('is_valid', True):
            skipped += 1
            continue
        name = (item.get('name') or '').strip()
        date_text = (item.get('date') or '').strip()
        if not name or not date_text:
            skipped += 1
            continue
        try:
            holiday_date = datetime.strptime(date_text, '%Y-%m-%d').date()
        except ValueError:
            errors.append(f'Invalid date for {name}.')
            skipped += 1
            continue

        holiday, was_created = CompanyHoliday.objects.update_or_create(
            date=holiday_date,
            defaults={
                'name': name,
                'holiday_type': normalize_holiday_type(item.get('holiday_type')),
                'description': (item.get('description') or '').strip(),
                'is_active': True,
            },
        )
        if was_created:
            created += 1
        else:
            updated += 1

    return {
        'created': created,
        'updated': updated,
        'skipped': skipped,
        'errors': errors,
    }
