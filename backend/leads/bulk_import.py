import re
import uuid
from datetime import date, datetime

from django.core.files.uploadedfile import UploadedFile
from django.db import transaction

from leads.models import Lead, LeadBulkUploadSession, LeadContact, LeadImportHistory

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
PREVIEW_ROW_LIMIT = 10
MULTIPLE_SHEETS_MESSAGE = 'Multiple sheets found. Please select one sheet to import.'

SYSTEM_FIELDS = {
    'company_name': {'label': 'Company Name', 'required': True},
    'website': {'label': 'Website', 'required': False},
    'country': {'label': 'Country', 'required': False},
    'industry': {'label': 'Industry', 'required': False},
    'company_size': {'label': 'Company Size', 'required': False},
    'service_fit': {'label': 'Service Fit', 'required': False},
    'current_status': {'label': 'Status', 'required': False},
    'source': {'label': 'Source', 'required': False},
    'priority': {'label': 'Priority', 'required': False},
    'remarks': {'label': 'Remarks', 'required': False},
    'next_follow_up_date': {'label': 'Next Follow-up Date', 'required': False},
    'decision_maker_name': {'label': 'Decision Maker Name', 'required': False},
    'decision_maker_designation': {'label': 'Decision Maker Designation', 'required': False},
    'decision_maker_linkedin': {'label': 'Decision Maker LinkedIn', 'required': False},
    'decision_maker_email': {'label': 'Decision Maker Email', 'required': False},
    'decision_maker_phone': {'label': 'Decision Maker Phone', 'required': False},
}

DEFAULT_COUNTRY = 'Not specified'
DEFAULT_INDUSTRY = 'General'
DEFAULT_SERVICE_FIT = Lead.ServiceFit.ANTRO_WORKFORCE
DEFAULT_STATUS = Lead.CurrentStatus.NEW
DEFAULT_PRIORITY = Lead.Priority.MEDIUM

SERVICE_FIT_ALIASES = {
    'antro workforce': Lead.ServiceFit.ANTRO_WORKFORCE,
    'antro workforce services': Lead.ServiceFit.ANTRO_WORKFORCE,
    'wodena technology': Lead.ServiceFit.WODENA_TECHNOLOGY,
    'wodena technology services': Lead.ServiceFit.WODENA_TECHNOLOGY,
    'igolo interior': Lead.ServiceFit.IGOLO_INTERIOR,
    'igolo interior services': Lead.ServiceFit.IGOLO_INTERIOR,
    'technology solutions': Lead.ServiceFit.TECHNOLOGY_SOLUTIONS,
}

STATUS_ALIASES = {choice.label.lower(): choice.value for choice in Lead.CurrentStatus}
STATUS_ALIASES.update({choice.value.lower(): choice.value for choice in Lead.CurrentStatus})

PRIORITY_ALIASES = {choice.label.lower(): choice.value for choice in Lead.Priority}
PRIORITY_ALIASES.update({choice.value.lower(): choice.value for choice in Lead.Priority})


def _cell_value(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _parse_xlsx(file_obj) -> dict:
    from openpyxl import load_workbook

    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    sheets = {}
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        rows = []
        for row in worksheet.iter_rows(values_only=True):
            rows.append([_cell_value(cell) for cell in row])
        if not rows:
            continue
        headers = rows[0]
        data_rows = [row for row in rows[1:] if any(cell for cell in row)]
        sheets[sheet_name] = {'headers': headers, 'rows': data_rows}
    workbook.close()
    return sheets


def _parse_xls(file_obj) -> dict:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=file_obj.read())
    sheets = {}
    for sheet in workbook.sheets():
        rows = []
        for row_idx in range(sheet.nrows):
            rows.append([_cell_value(sheet.cell_value(row_idx, col_idx)) for col_idx in range(sheet.ncols)])
        if not rows:
            continue
        headers = rows[0]
        data_rows = [row for row in rows[1:] if any(cell for cell in row)]
        sheets[sheet.name] = {'headers': headers, 'rows': data_rows}
    return sheets


def parse_lead_upload_file(uploaded_file: UploadedFile) -> dict:
    if uploaded_file.size > MAX_UPLOAD_BYTES:
        raise ValueError('File size must be 5 MB or less.')

    filename = (uploaded_file.name or '').lower()
    file_obj = uploaded_file.file
    file_obj.seek(0)

    if filename.endswith(('.xlsx', '.xlsm', '.xltx', '.xltm')):
        sheets = _parse_xlsx(file_obj)
    elif filename.endswith('.xls'):
        sheets = _parse_xls(file_obj)
    else:
        raise ValueError('Unsupported file type. Upload Excel (.xlsx or .xls).')

    if not sheets:
        raise ValueError('The uploaded Excel file has no readable sheets.')

    return sheets


def create_upload_session(user, uploaded_file: UploadedFile) -> LeadBulkUploadSession:
    sheets = parse_lead_upload_file(uploaded_file)
    return LeadBulkUploadSession.objects.create(
        id=uuid.uuid4(),
        file_name=uploaded_file.name,
        uploaded_by=user,
        sheets_data=sheets,
    )


def build_upload_parse_response(session: LeadBulkUploadSession) -> dict:
    sheets = []
    for sheet_name, sheet_data in session.sheets_data.items():
        headers = sheet_data.get('headers', [])
        rows = sheet_data.get('rows', [])
        preview_rows = []
        for index, row in enumerate(rows[:PREVIEW_ROW_LIMIT]):
            preview_rows.append({
                'row_number': index + 2,
                'values': row,
            })
        sheets.append({
            'name': sheet_name,
            'headers': headers,
            'preview_rows': preview_rows,
            'total_rows': len(rows),
        })

    sheet_count = len(sheets)
    return {
        'upload_id': str(session.id),
        'file_name': session.file_name,
        'system_fields': [
            {'key': key, 'label': meta['label'], 'required': meta['required']}
            for key, meta in SYSTEM_FIELDS.items()
        ],
        'sheets': sheets,
        'sheet_count': sheet_count,
        'requires_sheet_selection': sheet_count > 1,
        'multiple_sheets_message': MULTIPLE_SHEETS_MESSAGE if sheet_count > 1 else '',
    }


def _normalize_key(value: str) -> str:
    return re.sub(r'\s+', ' ', (value or '').strip().lower())


def _normalize_phone(value: str) -> str:
    return re.sub(r'\D', '', value or '')


def _normalize_website(value: str) -> str:
    website = (value or '').strip()
    if not website:
        return ''
    if not website.startswith(('http://', 'https://')):
        website = f'https://{website}'
    return website


def _parse_service_fit(value: str) -> str:
    raw = (value or '').strip()
    if not raw:
        return DEFAULT_SERVICE_FIT
    upper = raw.upper().replace(' ', '_')
    if upper in Lead.ServiceFit.values:
        return upper
    alias = SERVICE_FIT_ALIASES.get(raw.lower())
    if alias:
        return alias
    return DEFAULT_SERVICE_FIT


def _parse_status(value: str) -> str:
    raw = (value or '').strip()
    if not raw:
        return DEFAULT_STATUS
    alias = STATUS_ALIASES.get(raw.lower())
    if alias:
        return alias
    normalized = raw.upper().replace(' ', '_').replace('-', '_')
    if normalized in Lead.CurrentStatus.values:
        return normalized
    return DEFAULT_STATUS


def _parse_priority(value: str) -> str:
    raw = (value or '').strip()
    if not raw:
        return DEFAULT_PRIORITY
    alias = PRIORITY_ALIASES.get(raw.lower())
    if alias:
        return alias
    upper = raw.upper()
    if upper in Lead.Priority.values:
        return upper
    return DEFAULT_PRIORITY


def _parse_date(value: str):
    raw = (value or '').strip()
    if not raw:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%m/%d/%Y', '%d %b %Y', '%d-%b-%Y'):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None


def _row_dict_from_mapping(headers: list, row: list, mapping: dict) -> dict:
    header_index = {header: idx for idx, header in enumerate(headers) if header}
    mapped = {}
    for field_key, excel_column in mapping.items():
        if not excel_column:
            continue
        column_index = header_index.get(excel_column)
        if column_index is None:
            mapped[field_key] = ''
            continue
        mapped[field_key] = row[column_index] if column_index < len(row) else ''
    return mapped


def _duplicate_key(company_name: str, website: str) -> tuple:
    return (_normalize_key(company_name), _normalize_key(_normalize_website(website)))


def _existing_duplicate_keys(user) -> set:
    queryset = Lead.objects.filter(lead_owner=user)
    keys = set()
    for lead in queryset.only('company_name', 'website'):
        keys.add(_duplicate_key(lead.company_name, lead.website))
    return keys


def _contact_identity_key(email: str, phone: str, linkedin: str) -> tuple:
    return (
        _normalize_key(email),
        _normalize_phone(phone),
        _normalize_key(_normalize_website(linkedin)),
    )


def _validate_decision_maker(mapped: dict, issues: list) -> None:
    name = (mapped.get('decision_maker_name') or '').strip()
    if not name:
        return
    email = (mapped.get('decision_maker_email') or '').strip()
    phone = (mapped.get('decision_maker_phone') or '').strip()
    linkedin = (mapped.get('decision_maker_linkedin') or '').strip()
    if not any([email, phone, linkedin]):
        issues.append('Decision maker requires Email, Phone, or LinkedIn')


def _contact_duplicate_on_lead(lead: Lead, email: str, phone: str, linkedin: str) -> bool:
    email_key, phone_key, linkedin_key = _contact_identity_key(email, phone, linkedin)
    if not any([email_key, phone_key, linkedin_key]):
        return False
    for contact in lead.contacts.all():
        contact_email, contact_phone, contact_linkedin = _contact_identity_key(
            contact.email,
            contact.phone,
            contact.linkedin_profile_url,
        )
        if email_key and contact_email == email_key:
            return True
        if phone_key and contact_phone == phone_key:
            return True
        if linkedin_key and contact_linkedin == linkedin_key:
            return True
    return False


def _build_preview_row(row_number: int, mapped: dict, existing_keys: set, seen_keys: set, seen_contacts: set) -> dict:
    company_name = (mapped.get('company_name') or '').strip()
    website = _normalize_website(mapped.get('website', ''))
    issues = []
    is_duplicate = False

    if not company_name:
        issues.append('Missing company name')
    else:
        key = _duplicate_key(company_name, website)
        if key in existing_keys:
            is_duplicate = True
            issues.append('Duplicate of existing lead (same company, website, and owner)')
        elif key in seen_keys:
            is_duplicate = True
            issues.append('Duplicate within file (same company, website, and owner)')
        else:
            seen_keys.add(key)

    _validate_decision_maker(mapped, issues)

    dm_name = (mapped.get('decision_maker_name') or '').strip()
    dm_email = (mapped.get('decision_maker_email') or '').strip()
    dm_phone = (mapped.get('decision_maker_phone') or '').strip()
    dm_linkedin = _normalize_website(mapped.get('decision_maker_linkedin', ''))
    if dm_name and any([dm_email, dm_phone, dm_linkedin]):
        contact_key = _contact_identity_key(dm_email, dm_phone, dm_linkedin)
        if any(contact_key):
            if contact_key in seen_contacts:
                issues.append('Duplicate decision maker contact in file (same Email, Phone, or LinkedIn)')
            else:
                seen_contacts.add(contact_key)

    next_follow_up = _parse_date(mapped.get('next_follow_up_date', ''))
    if (mapped.get('next_follow_up_date') or '').strip() and next_follow_up is None:
        issues.append('Invalid next follow-up date format')

    blocking_issues = [issue for issue in issues if not issue.startswith('Duplicate of existing lead') and not issue.startswith('Duplicate within file')]

    return {
        'row_number': row_number,
        'company_name': company_name,
        'website': website,
        'country': (mapped.get('country') or '').strip() or DEFAULT_COUNTRY,
        'industry': (mapped.get('industry') or '').strip() or DEFAULT_INDUSTRY,
        'company_size': (mapped.get('company_size') or '').strip(),
        'source': (mapped.get('source') or '').strip(),
        'service_fit': _parse_service_fit(mapped.get('service_fit', '')),
        'current_status': _parse_status(mapped.get('current_status', '')),
        'priority': _parse_priority(mapped.get('priority', '')),
        'remarks': (mapped.get('remarks') or '').strip(),
        'next_follow_up_date': next_follow_up.isoformat() if next_follow_up else '',
        'decision_maker_name': dm_name,
        'decision_maker_designation': (mapped.get('decision_maker_designation') or '').strip(),
        'decision_maker_linkedin': dm_linkedin,
        'decision_maker_email': dm_email,
        'decision_maker_phone': dm_phone,
        'issues': issues,
        'is_duplicate': is_duplicate,
        'can_import': bool(company_name) and not blocking_issues,
    }


def build_import_preview(session: LeadBulkUploadSession, sheet_name: str, mapping: dict) -> dict:
    sheet_data = session.sheets_data.get(sheet_name)
    if not sheet_data:
        raise ValueError('Selected sheet was not found in the uploaded file.')

    headers = sheet_data.get('headers', [])
    rows = sheet_data.get('rows', [])
    existing_keys = _existing_duplicate_keys(session.uploaded_by)
    seen_keys = set()
    seen_contacts = set()

    preview_rows = []
    missing_company_count = 0
    duplicate_count = 0
    contact_issue_count = 0

    for index, row in enumerate(rows):
        row_number = index + 2
        mapped = _row_dict_from_mapping(headers, row, mapping)
        preview_row = _build_preview_row(row_number, mapped, existing_keys, seen_keys, seen_contacts)
        preview_rows.append(preview_row)

        if not preview_row['company_name']:
            missing_company_count += 1
        if preview_row['is_duplicate']:
            duplicate_count += 1
        if any('Decision maker' in issue or 'decision maker' in issue.lower() for issue in preview_row['issues']):
            contact_issue_count += 1

    importable_if_skip_duplicates = sum(
        1 for row in preview_rows if row['can_import'] and not row['is_duplicate']
    )

    return {
        'upload_id': str(session.id),
        'file_name': session.file_name,
        'sheet_name': sheet_name,
        'mapping': mapping,
        'summary': {
            'total_rows': len(rows),
            'valid_rows': sum(1 for row in preview_rows if row['can_import']),
            'importable_rows': importable_if_skip_duplicates,
            'missing_company_rows': missing_company_count,
            'duplicate_rows': duplicate_count,
            'contact_issue_rows': contact_issue_count,
        },
        'preview_rows': preview_rows[:PREVIEW_ROW_LIMIT],
        'all_rows': preview_rows,
    }


def execute_import(
    session: LeadBulkUploadSession,
    sheet_name: str,
    mapping: dict,
    skip_duplicates: bool = True,
) -> dict:
    preview = build_import_preview(session, sheet_name, mapping)
    user = session.uploaded_by

    imported_rows = 0
    skipped_rows = 0
    failed_rows = 0
    error_details = []

    with transaction.atomic():
        for row in preview['all_rows']:
            row_errors = list(row['issues'])

            if not row['company_name']:
                failed_rows += 1
                error_details.append({'row_number': row['row_number'], 'errors': row_errors or ['Missing company name']})
                continue

            if row['is_duplicate'] and skip_duplicates:
                skipped_rows += 1
                error_details.append({
                    'row_number': row['row_number'],
                    'errors': ['Skipped duplicate lead'],
                    'status': 'skipped',
                })
                continue

            if not row['can_import']:
                failed_rows += 1
                error_details.append({'row_number': row['row_number'], 'errors': row_errors})
                continue

            try:
                lead = Lead.objects.create(
                    company_name=row['company_name'],
                    website=row['website'],
                    country=row['country'],
                    industry=row['industry'],
                    company_size=row['company_size'],
                    source=row['source'],
                    service_fit=row['service_fit'],
                    current_status=row['current_status'],
                    priority=row['priority'],
                    remarks=row['remarks'],
                    next_follow_up_date=row['next_follow_up_date'] or None,
                    created_by=user,
                    lead_owner=user,
                    last_updated_by=user,
                )

                if row['decision_maker_name']:
                    if _contact_duplicate_on_lead(
                        lead,
                        row['decision_maker_email'],
                        row['decision_maker_phone'],
                        row['decision_maker_linkedin'],
                    ):
                        error_details.append({
                            'row_number': row['row_number'],
                            'errors': ['Duplicate decision maker contact on lead (not added)'],
                            'status': 'warning',
                        })
                    else:
                        LeadContact.objects.create(
                            lead=lead,
                            full_name=row['decision_maker_name'],
                            designation=row['decision_maker_designation'] or 'Not specified',
                            linkedin_profile_url=row['decision_maker_linkedin'],
                            email=row['decision_maker_email'],
                            phone=row['decision_maker_phone'],
                        )
                imported_rows += 1
            except Exception as exc:
                failed_rows += 1
                error_details.append({
                    'row_number': row['row_number'],
                    'errors': [str(exc)],
                    'status': 'failed',
                })

        history = LeadImportHistory.objects.create(
            file_name=session.file_name,
            uploaded_by=user,
            total_rows=preview['summary']['total_rows'],
            imported_rows=imported_rows,
            skipped_rows=skipped_rows,
            failed_rows=failed_rows,
            sheet_name=sheet_name,
            column_mapping=mapping,
            error_details=error_details,
        )

        session.delete()

    from leads.risk_events import record_bulk_upload_duplicates_event

    record_bulk_upload_duplicates_event(user, skipped_rows, history.file_name)

    return {
        'import_id': history.id,
        'file_name': history.file_name,
        'sheet_name': sheet_name,
        'total_rows': history.total_rows,
        'imported_rows': history.imported_rows,
        'skipped_rows': history.skipped_rows,
        'failed_rows': history.failed_rows,
        'error_details': error_details,
    }


def list_import_history(user):
    queryset = LeadImportHistory.objects.select_related('uploaded_by')
    if not user.is_super_admin:
        queryset = queryset.filter(uploaded_by=user)
    return queryset
